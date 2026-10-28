"""얼리액세스 기능흐름도 생성: 템플릿 머리글 양식 유지, 예시 슬라이드는 제거, 흐름도는 편집 가능한 도형으로."""
import copy, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

SRC, OUT = sys.argv[1], sys.argv[2]
TEAM = "얼리액세스"
FONT = "맑은 고딕"
C = dict(navy="1F3864", blue="4A6FA5", ink="262626", gray="5F6368", line="4D4D4D",
         term_f="E6E6E6", term_l="8C8C8C", start_f="ECE6FA", start_l="8E7CC3",
         proc_f="EEF3FB", io_f="FFF4D6", io_l="C99A06",
         dec_f="EAF4E4", dec_l="5E9E4A", db_f="E3F1FA", page_f="F7F9FC", err="C0392B", err_f="FDEDEB",
         note_f="F4F7FB", ok="2E7D32", warn="D98200", orange="E07B39", green="2E7D4F", slate="5B6B82")
ICONS = __import__('pathlib').Path(__file__).parent / 'flowchart_icons'
DY = 0.0  # 흐름도 쪽은 안쪽 제목 띠만큼 아래로 내림
SQ = False  # True: 흐름도를 띠 아래 ~ 하단 '핵심 포인트' 줄 위로 세로 압축
SQ_Y0, SQ_L0, SQ_K = 0.9, 0.55, 0.84


def Y(y):
    return SQ_Y0 + (y - SQ_L0) * SQ_K if SQ else y + DY


def H(h):
    return h * SQ_K if SQ else h
rgb = lambda h: RGBColor.from_string(h)

prs = Presentation(SRC)
tmpl = prs.slides[1]
hdr_title, hdr_team = [copy.deepcopy(s._element) for s in list(tmpl.shapes)[:2]]
layout = tmpl.slide_layout


def set_font(run, size, bold=False, color="262626"):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, FONT
    f.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def put_text(tf, lines, size=8, bold=False, color="262626", align=PP_ALIGN.CENTER):
    """lines: str 또는 [(text, {size,bold,color})] / 줄바꿈은 \n"""
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    items = lines if isinstance(lines, list) else [(l, {}) for l in lines.split("\n")]
    for i, (t, o) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = o.get("align", align)
        r = p.add_run()
        r.text = t
        set_font(r, o.get("size", size), o.get("bold", bold), o.get("color", color))


def new_slide(title):
    s = prs.slides.add_slide(layout)
    for el, txt in ((hdr_title, title), (hdr_team, TEAM)):
        e = copy.deepcopy(el)
        s.shapes._spTree.append(e)
        runs = e.findall(".//" + qn("a:r"))
        runs[0].find(qn("a:t")).text = txt
        if el is hdr_title:  # 템플릿 제목칸 줄간격 154% → 글자가 머리글 띠 밖(흰 배경)으로 내려가 잘림. 100%로
            e.find(".//" + qn("a:lnSpc") + "/" + qn("a:spcPct")).set("val", "100000")
        for r in runs[1:]:
            r.getparent().remove(r)
    return s


KIND = {
    "term": (MSO_SHAPE.FLOWCHART_TERMINATOR, "term_f", "term_l"),
    "proc": (MSO_SHAPE.ROUNDED_RECTANGLE, "proc_f", "blue"),
    "io": (MSO_SHAPE.FLOWCHART_DATA, "io_f", "io_l"),
    "dec": (MSO_SHAPE.FLOWCHART_DECISION, "dec_f", "dec_l"),
    "db": (MSO_SHAPE.FLOWCHART_MAGNETIC_DISK, "db_f", "blue"),
    "page": (MSO_SHAPE.FLOWCHART_DOCUMENT, "page_f", "blue"),
    "err": (MSO_SHAPE.ROUNDED_RECTANGLE, "err_f", "err"),
    "ok": (MSO_SHAPE.ROUNDED_RECTANGLE, "proc_f", "ok"),
}


def node(s, kind, x, y, w, h, text, size=7.5, bold=False, dash=False):
    shp, f, l = KIND[kind]
    if kind == "term" and text.startswith("시작"):
        f, l = "start_f", "start_l"
    sh = s.shapes.add_shape(shp, Inches(x), Inches(Y(y)), Inches(w), Inches(H(h)))
    if shp == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = 0.18
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(C[f])
    sh.line.color.rgb = rgb(C[l]); sh.line.width = Pt(1.1)
    if dash:
        sh.line.dash_style = 4  # dash
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    color = C["err"] if kind == "err" else C["ink"]
    put_text(tf, text, size=size * (0.94 if SQ else 1), bold=bold or kind == "term", color=color)
    sh._box = (x, y, w, h)
    return sh


def edge(s, pts, label=None, lpos=None, color="line", dashed=False, arrow=True):
    """꺾인 선: pts = [(x,y),...] 인치. 마지막 구간에만 화살촉."""
    for i in range(len(pts) - 1):
        (x1, y1), (x2, y2) = pts[i], pts[i + 1]
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(Y(y1)), Inches(x2), Inches(Y(y2)))
        c.line.color.rgb = rgb(C[color]); c.line.width = Pt(1)
        ln = c.line._get_or_add_ln()
        if dashed:
            etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
        if arrow and i == len(pts) - 2:
            t = etree.SubElement(ln, qn("a:tailEnd")); t.set("type", "triangle"); t.set("w", "med"); t.set("len", "med")
    if label:
        lx, ly = lpos
        tb = s.shapes.add_textbox(Inches(lx), Inches(Y(ly)), Inches(0.22), Inches(0.16))
        put_text(tb.text_frame, label, size=7, bold=True, color=C["blue"] if label == "Y" else C["err"])
        tb.text_frame.margin_left = tb.text_frame.margin_right = 0


def text(s, x, y, w, h, lines, size=8, align=PP_ALIGN.LEFT, fill=None, line=None, anchor=MSO_ANCHOR.TOP):
    if fill or line:
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(Y(y)), Inches(w), Inches(H(h)))
        sh.adjustments[0] = 0.06
        sh.shadow.inherit = False
        if fill:
            sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = rgb(line); sh.line.width = Pt(0.8)
        else:
            sh.line.fill.background()
        tf = sh.text_frame
    else:
        sh = s.shapes.add_textbox(Inches(x), Inches(Y(y)), Inches(w), Inches(H(h)))
        tf = sh.text_frame
    tf.vertical_anchor = anchor
    put_text(tf, lines, size=size, align=align)
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    return sh


# 좌표 도우미
def cx(n): x, y, w, h = n._box; return x + w / 2
def cy(n): x, y, w, h = n._box; return y + h / 2
def top(n): x, y, w, h = n._box; return (x + w / 2, y)
def bot(n): x, y, w, h = n._box; return (x + w / 2, y + h)
def lft(n): x, y, w, h = n._box; return (x, y + h / 2)
def rgt(n): x, y, w, h = n._box; return (x + w, y + h / 2)
def down(s, a, b, **k): edge(s, [bot(a), top(b)], **k)
def up(s, a, b, **k): edge(s, [top(a), bot(b)], **k)


def icon(s, name, x, y, sz, white=False):
    s.shapes.add_picture(f"{ICONS}/{name}_{'w' if white else 'n'}.png", Inches(x), Inches(Y(y)), Inches(sz), Inches(sz))


def note_card(s, x, y, w, h, title, bullets, size=7.5, ic="MdInfoOutline"):
    h = h if SQ else h - DY  # 제목 띠만큼 내려간 쪽은 카드 높이를 줄여 로고와 안 겹치게
    lines = [("      " + title, {"bold": True, "size": size + 1, "color": C["navy"]})] + [("• " + b, {"size": size}) for b in bullets]
    sh = text(s, x, y, w, h, lines, fill=C["note_f"], line="C9D3E6")
    sh.text_frame.paragraphs[0].space_after = Pt(5)
    icon(s, ic, x + 0.09, y + 0.06, 0.17)
    return sh


def points(s, items, title="핵심\n포인트"):
    y0, hh = 4.6, 0.5
    st = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.3), Inches(y0), Inches(9.4), Inches(hh))
    st.adjustments[0] = 0.15; st.shadow.inherit = False
    st.fill.solid(); st.fill.fore_color.rgb = rgb(C["note_f"]); st.line.color.rgb = rgb("C9D3E6")
    lb = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.3), Inches(y0), Inches(0.85), Inches(hh))
    lb.adjustments[0] = 0.15; lb.shadow.inherit = False
    lb.fill.solid(); lb.fill.fore_color.rgb = rgb(C["navy"]); lb.line.fill.background()
    lb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    put_text(lb.text_frame, [(t, {"bold": True, "size": 7.5, "color": "FFFFFF"}) for t in title.split("\n")])
    step = 8.4 / len(items)
    for i, (ic, t1, t2) in enumerate(items):
        x = 1.3 + i * step
        s.shapes.add_picture(f"{ICONS}/{ic}_n.png", Inches(x), Inches(y0 + 0.11), Inches(0.28), Inches(0.28))
        tb = s.shapes.add_textbox(Inches(x + 0.3), Inches(y0), Inches(step - 0.35), Inches(hh))
        tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        put_text(tb.text_frame, [(t1, {"bold": True, "size": 6.8}), (t2, {"size": 5.8, "color": C["gray"]})], align=PP_ALIGN.LEFT)


def banner(s, sub, ic="MdDirectionsBus"):
    """템플릿 예시(Rento)의 안쪽 제목 띠: 왼쪽 서비스명, 오른쪽 이 쪽의 흐름 이름"""
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.3), Inches(0.5), Inches(9.4), Inches(0.28))
    b.adjustments[0] = 0.25; b.shadow.inherit = False
    b.fill.solid(); b.fill.fore_color.rgb = rgb(C["navy"]); b.line.fill.background()
    s.shapes.add_picture(f"{ICONS}/{ic}_w.png", Inches(0.42), Inches(0.535), Inches(0.21), Inches(0.21))
    for x, w, t, al, sz in ((0.68, 3.5, [("한자리", {"bold": True, "size": 10, "color": "FFFFFF"}), ], PP_ALIGN.LEFT, 10),
                            (5.2, 4.4, [(sub, {"bold": True, "size": 9, "color": "FFFFFF"})], PP_ALIGN.RIGHT, 9)):
        tb = s.shapes.add_textbox(Inches(x), Inches(0.5), Inches(w), Inches(0.28))
        tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        put_text(tb.text_frame, t, align=al)
        tb.text_frame.margin_top = tb.text_frame.margin_bottom = 0
    tb = s.shapes.add_textbox(Inches(1.35), Inches(0.5), Inches(3.0), Inches(0.28))
    tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    put_text(tb.text_frame, [("첫 버스, 탈 수 있을까?", {"size": 7, "color": "C9D3E6"})], align=PP_ALIGN.LEFT)
    tb.text_frame.margin_top = tb.text_frame.margin_bottom = 0


# ───────────────────────── 1. 범례 및 개요
s = new_slide("기능 흐름도 - 범례 및 개요")
text(s, 0.35, 0.55, 4.5, 0.3, [("흐름도 기호", {"bold": True, "size": 11, "color": C["navy"]})])
rows = [("term", "시작 / 끝", "흐름의 시작과 끝"), ("proc", "처리", "시스템이 처리하는 단계"),
        ("io", "입력", "사용자가 값을 입력하는 단계"), ("dec", "판단", "조건에 따라 Y / N 으로 나뉨"),
        ("page", "화면", "사용자에게 보여주는 화면"), ("db", "데이터·외부 API", "저장소, 버스 API, 모델"),
        ("err", "예외 처리", "오류·보류 등 예외 안내"), (None, "연결선", "진행 방향 (Y: 예, N: 아니오)")]
y = 0.92
for k, name, desc in rows:
    if k:
        node(s, k, 0.45, y, 1.05, 0.36, name if k != "dec" else "판단", size=7)
    else:
        edge(s, [(0.5, y + 0.18), (1.45, y + 0.18)])
    text(s, 1.65, y + 0.03, 0.95, 0.3, [(name, {"bold": True, "size": 8.5})])
    text(s, 2.55, y + 0.03, 2.3, 0.3, [(desc, {"size": 8, "color": C["gray"]})])
    y += 0.49

text(s, 5.15, 0.55, 4.5, 0.3, [("서비스 개요", {"bold": True, "size": 11, "color": C["navy"]})])
text(s, 5.15, 0.9, 4.5, 1.1, [
    ("한자리 — 조회 시점의 첫 번째 버스에 탈 수 있는지 알려주는 반응형 웹", {"bold": True, "size": 8.5}),
    ("대상: 등교 방향 6개 정류장", {"size": 8, "bold": True, "color": C["navy"]}),
    ("• 성북02  한성대입구역 5번출구 → 2번출구 → 녹십자약국", {"size": 8}),
    ("• 종로03  동대문역.흥인지문 → 동묘앞역 → 창신역4번출구", {"size": 8}),
    ("입력: 내 앞 대기 인원 1개 / 출력: 여유 · 빠듯 · 어려움 + 근거", {"size": 8}),
], fill=C["note_f"], line="C9D3E6")
tbl = s.shapes.add_table(9, 3, Inches(5.15), Inches(2.1), Inches(4.5), Inches(2.7)).table
tbl.columns[0].width, tbl.columns[1].width, tbl.columns[2].width = Inches(0.6), Inches(2.6), Inches(1.3)
data = [("쪽", "흐름", "기능 ID (PRD)"), ("2", "전체 (단계 요약)", "F-01 ~ F-07"), ("3", "전체 상세 흐름", "F-01 ~ F-07"),
        ("4", "도착정보 조회 · 상태 안내", "F-01, F-05"), ("5", "탑승 가능성 예측 · 표시", "F-02 ~ F-04"),
        ("6", "부록 A  데이터 수집 · 학습 · 평가", "F-06, F-07"), ("7", "부록 B  AI 예측 모델 내부", "F-03"),
        ("8", "부록 C  검증 현황 · 제안", "F-07")]
for r, row in enumerate(data):
    for c_, v in enumerate(row):
        cell = tbl.cell(r, c_)
        cell.text = ""
        put_text(cell.text_frame, [(v, {"bold": r == 0, "size": 8.5, "color": "FFFFFF" if r == 0 else C["ink"]})],
                 align=PP_ALIGN.CENTER if c_ != 1 else PP_ALIGN.LEFT)
        cell.fill.solid(); cell.fill.fore_color.rgb = rgb(C["navy"] if r == 0 else ("FFFFFF" if r % 2 else "F2F5FA"))
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
text(s, 5.15, 4.86, 4.5, 0.22, [("v1.0 · 2026-09-28 · 작성: 최지용 (AI·데이터) · 기준 문서: PRD v0.9, WBS v1.0", {"size": 6.5, "color": C["gray"]})])

# ───────────────────────── 2. 전체 서비스 흐름

# ───────────────────────── 2. 전체 (단계 요약, 템플릿 예시 2쪽 형식)
DY = 0.0
s = new_slide("기능 흐름도 - 전체"); banner(s, "기능 흐름도 - 전체")
text(s, 0.3, 0.84, 6.0, 0.3, [("기능 흐름 및 의사결정 프로세스 (전체)", {"bold": True, "size": 11, "color": C["navy"]})])
text(s, 0.3, 1.1, 6.0, 0.25, [("□ 한자리의 전체 서비스 흐름을 나타냅니다.", {"size": 8})])
steps = [
    ("MdPlace", "정류장 선택", "", "navy", ["노선: 성북02 · 종로03", "정류장: 등교 방향 6곳"]),
    ("MdDirectionsBus", "도착정보 확인", "F-01", "navy", ["첫째 · 둘째 버스 도착", "남은 정류장 · 위치", "갱신 시각 표시"]),
    ("MdPeople", "대기 인원 입력", "F-02", "navy", ["내 앞 대기 인원 입력", "0~40명 정수", "상황이 바뀌면 수정"]),
    ("MdInsights", "탑승 가능성 예측", "F-03", "green", ["지금 버스 재차 확인", "남은 정류장 승차 추정", "도착 시 인원 → 확률"]),
    ("MdTaskAlt", "결과 확인", "F-04", "green", ["여유 · 빠듯 · 어려움", "근거: 재차 · 몇 정류장 전", "불확실하면 '판단 일러요'"]),
    ("MdTune", "대기 여부 판단", "", "green", ["계속 기다리기", "둘째 버스 · 다른 방법", "결정은 사용자가"]),
    ("MdRefresh", "재조회 · 기록", "F-06", "slate", ["30초마다 자동 갱신", "버스 바뀌면 재예측", "(선택) 탑승 결과 기록"]),
]
W0, G0, X0, Y0, H0 = 1.1, 0.24, 0.35, 1.98, 0.9
xs = [X0 + i * (W0 + G0) for i in range(len(steps))]
for i, (ic, name, fid, col, bl) in enumerate(steps):
    x = xs[i]
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(Y0), Inches(W0), Inches(H0))
    card.adjustments[0] = 0.12; card.shadow.inherit = False
    card.fill.solid(); card.fill.fore_color.rgb = rgb(C[col]); card.line.fill.background()
    icon(s, ic, x + W0 / 2 - 0.17, Y0 + 0.1, 0.34, white=True)
    tb = s.shapes.add_textbox(Inches(x), Inches(Y0 + 0.47), Inches(W0), Inches(0.4))
    put_text(tb.text_frame, [(name, {"bold": True, "size": 8, "color": "FFFFFF"})] +
             ([(f"({fid})", {"size": 6.5, "color": "DDE5F2"})] if fid else []))
    tb.text_frame.margin_top = 0
    if i < len(steps) - 1:
        edge(s, [(x + W0 + 0.02, Y0 + H0 / 2), (xs[i + 1] - 0.02, Y0 + H0 / 2)])
    text(s, x - 0.1, Y0 + H0 + 0.06, W0 + 0.22, 0.75, [("• " + b, {"size": 6.3}) for b in bl])
# 예외 · 의사결정 처리 (점선 주황)
for i, t1, t2 in ((1, "상태 안내 (F-05)", "운행 종료 · API 장애 · 정보 없음"), (4, "확률 보류 (F-04)", "검증 안 된 구간 · 정보 60초 이상")):
    x = xs[i] - 0.35
    bx = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.33), Inches(W0 + 0.7), Inches(0.44))
    bx.adjustments[0] = 0.15; bx.shadow.inherit = False
    bx.fill.solid(); bx.fill.fore_color.rgb = rgb("FFF4EC")
    bx.line.color.rgb = rgb(C["orange"]); bx.line.width = Pt(1.1); bx.line.dash_style = 4
    icon(s, "MdWarningAmber", x + 0.08, 1.43, 0.22)
    tb = s.shapes.add_textbox(Inches(x + 0.3), Inches(1.33), Inches(W0 + 0.38), Inches(0.44))
    tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    put_text(tb.text_frame, [(t1, {"bold": True, "size": 7, "color": C["orange"]}), (t2, {"size": 6.2, "color": C["gray"]})], align=PP_ALIGN.LEFT)
    edge(s, [(xs[i] + W0 / 2, 1.77), (xs[i] + W0 / 2, Y0 - 0.01)], color="orange", dashed=True)
# 범례 줄
lx = 0.35
for col, lab, dashed in (("navy", "정보 확인 단계", False), ("green", "예측 · 판단 단계", False), ("slate", "갱신 · 기록", False), ("orange", "예외 · 의사결정 처리", True)):
    sq = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(lx), Inches(3.72), Inches(0.14), Inches(0.14))
    sq.shadow.inherit = False
    if dashed:
        sq.fill.background(); sq.line.color.rgb = rgb(C[col]); sq.line.dash_style = 4
    else:
        sq.fill.solid(); sq.fill.fore_color.rgb = rgb(C[col]); sq.line.fill.background()
    text(s, lx + 0.14, 3.67, 1.5, 0.24, [(lab, {"size": 6.5, "color": C["gray"]})])
    lx += 1.55
# 공통 지원 기능
strip = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.3), Inches(4.08), Inches(9.4), Inches(0.78))
strip.adjustments[0] = 0.12; strip.shadow.inherit = False
strip.fill.solid(); strip.fill.fore_color.rgb = rgb(C["note_f"]); strip.line.color.rgb = rgb("C9D3E6")
lab = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.3), Inches(4.08), Inches(1.0), Inches(0.78))
lab.adjustments[0] = 0.12; lab.shadow.inherit = False
lab.fill.solid(); lab.fill.fore_color.rgb = rgb(C["navy"]); lab.line.fill.background()
lab.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
put_text(lab.text_frame, [("공통 지원", {"bold": True, "size": 8.5, "color": "FFFFFF"}), ("기능", {"bold": True, "size": 8.5, "color": "FFFFFF"})])
commons = [("MdRefresh", "자동 갱신", "30초마다 · 수동 재조회"), ("MdNotificationsActive", "상태 안내", "운행 종료 · 장애 · 정보 없음"),
           ("MdVerified", "검증된 확률만", "기준 모델 통과 구간"), ("MdLock", "개인정보 미저장", "익명 · API 키 서버 보관"),
           ("MdEditNote", "평가 기록", "(선택) 실제 탑승 결과")]
for i, (ic, t1, t2) in enumerate(commons):
    x = 1.45 + i * 1.65
    icon(s, ic, x, 4.29, 0.36)
    text(s, x + 0.38, 4.2, 1.3, 0.55, [(t1, {"bold": True, "size": 7.2}), (t2, {"size": 6, "color": C["gray"]})], anchor=MSO_ANCHOR.MIDDLE)

s = new_slide("기능 흐름도 - 전체 상세"); SQ = True; banner(s, "전체 상세 흐름 (의사결정 · 예외 포함)")
points(s, [("MdRefresh", "30초 자동 재조회", "장애 · 운행 종료 때도"), ("MdPeople", "인원 바뀌면 재예측", "앞 대기 인원 수정"),
           ("MdVerified", "검증된 구간만 확률", "그 외는 도착정보만"), ("MdEditNote", "탑승 결과 기록", "(선택) 익명"),
           ("MdLock", "개인정보 미저장", "로그인 없이 사용")])
A, W = 0.4, 1.8          # 주 흐름 1열 (위→아래)
B, WB = 2.55, 1.45       # 1열 예외
Cc, WC = 4.3, 1.8        # 주 흐름 2열 (아래→위)
D, WD = 6.4, 1.45        # 2열 예외·데이터
n1 = node(s, "term", A + 0.15, 0.6, W - 0.3, 0.3, "시작 (한자리 접속)")
n2 = node(s, "page", A, 1.12, W, 0.42, "노선 · 정류장 선택 화면\n(성북02 · 종로03, 6곳)")
n3 = node(s, "proc", A, 1.78, W, 0.42, "도착정보 조회 (F-01)\n첫째 · 둘째 버스")
n4 = node(s, "dec", A, 2.42, W, 0.62, "운행 중이고\n정보가 최신인가?")
n5 = node(s, "io", A, 3.27, W, 0.42, "앞 대기 인원 입력\n(F-02)")
n6 = node(s, "dec", A, 3.92, W, 0.62, "입력이 유효한가?\n(정수 0~40명)")
e1 = node(s, "err", B, 2.5, WB, 0.46, "상태 안내 (F-05)\n종료 · 장애 · 정보 없음")
e2 = node(s, "err", B, 4.0, WB, 0.46, "입력 수정 안내\n(숫자 0~40 입력)")
db1 = node(s, "db", B + 0.1, 1.08, WB - 0.2, 0.5, "서울시 버스\n도착 · 위치 API", size=7)
down(s, n1, n2); down(s, n2, n3); down(s, n3, n4); down(s, n4, n5, label="Y", lpos=(A + W / 2 + 0.05, 3.05))
down(s, n5, n6)
edge(s, [rgt(n4), lft(e1)], label="N", lpos=(A + W + 0.02, 2.55))
edge(s, [top(e1), (B + WB / 2, cy(n3) - 0.08)], arrow=False)  # 30초 뒤 자동 재조회 → 재조회 선에 합류
text(s, B + WB / 2 + 0.02, 2.2, 1.2, 0.25, [("30초 뒤 재조회", {"size": 6.5, "color": C["gray"]})])
edge(s, [rgt(n6), lft(e2)], label="N", lpos=(A + W + 0.02, 4.05))
edge(s, [top(e2), (B + WB / 2, cy(n5)), rgt(n5)])
edge(s, [(cx(db1), 1.58), (cx(db1), 1.7), (A + W + 0.02, 1.7), (A + W + 0.02, 1.9)], color="blue", dashed=True)
m7 = node(s, "proc", Cc, 4.12, WC, 0.42, "탑승 가능성 예측 (F-03)\n과거 평균 + 오늘 실측 → 확률")
m8 = node(s, "dec", Cc, 3.3, WC, 0.62, "표시 조건 충족?\n(F-04)")
m9 = node(s, "page", Cc, 2.52, WC, 0.52, "결과 화면\n여유 · 빠듯 · 어려움 + 근거\n둘째 버스 · 갱신 시각")
m10 = node(s, "dec", Cc, 1.7, WC, 0.62, "재조회가\n필요한가?")
m11 = node(s, "io", Cc, 1.12, WC, 0.36, "탑승 결과 기록\n(F-06, 선택)")
m12 = node(s, "term", Cc + 0.35, 0.6, WC - 0.7, 0.3, "끝")
edge(s, [bot(n6), (cx(n6), 4.82), (cx(m7), 4.82), bot(m7)], label="Y", lpos=(cx(n6) + 0.05, 4.57))
up(s, m7, m8); up(s, m8, m9, label="Y", lpos=(cx(m8) + 0.05, 3.08)); up(s, m9, m10)
up(s, m10, m11, label="N", lpos=(cx(m10) + 0.05, 1.48)); up(s, m11, m12)
edge(s, [lft(m10), (A + W + 0.15, cy(m10)), (A + W + 0.15, cy(n3) - 0.08), (A + W, cy(n3) - 0.08)],
     label="Y", lpos=(Cc - 0.25, cy(m10) - 0.2))
e3 = node(s, "err", D, 3.38, WD, 0.46, "확률 보류\n도착정보만 표시")
edge(s, [rgt(m8), lft(e3)], label="N", lpos=(Cc + WC + 0.02, 3.42))
edge(s, [top(e3), (D + WD / 2, cy(m10)), rgt(m10)])
db2 = node(s, "db", D + 0.1, 4.08, WD - 0.2, 0.5, "예측 모델\n(규칙 + XGBoost)", size=7)
edge(s, [lft(db2), rgt(m7)], color="blue", dashed=True)
db3 = node(s, "db", D + 0.1, 1.04, WD - 0.2, 0.5, "평가 기록\n(익명)", size=7)
edge(s, [rgt(m11), lft(db3)], color="blue", dashed=True)
note_card(s, 8.05, 0.6, 1.6, 4.25, "핵심 규칙", ic="MdTaskAlt", bullets=[
    "버스가 멀어도 첫 버스 기준으로 늘 표시",
    "표시는 3단계 + 근거 (지금 재차 · 남은 정류장 · 갱신 시각)",
    "'여유'는 실제 탑승 97% 이상 검증된 구간에서만",
    "불확실하면 '판단하기 일러요'",
    "표시 조건: 검증 통과 구간 + 60초 내 정보",
    "대상 버스가 바뀌면 이전 예측 삭제",
    "개인정보 저장 안 함 (익명 기록)",
    "(제안) 입력된 대기 인원을 뒤 정류장 예측에 활용",
], size=7)

# ───────────────────────── 3. 도착정보 조회 · 상태 안내
s = new_slide("기능 흐름도 - 도착정보 조회 · 상태 안내"); banner(s, "정류장 조회 · 상태 안내 프로세스", "MdDirectionsBus")
points(s, [("MdTimer", "캐시 30초", "같은 정류장 재조회"), ("MdNotificationsActive", "상태 5종 구분", "정상 · 지연 · 장애 · 종료 · 없음"),
           ("MdEventAvailable", "운행 시간", "API 첫차 · 막차 시각"), ("MdRefresh", "30초 뒤 재시도", "장애 · 정보 없음"),
           ("MdLock", "API 키 서버 보관", "화면에 노출 안 함")])
A, W, B, WB = 0.4, 1.8, 2.5, 1.55
Cc, WC, D, WD = 4.3, 1.8, 6.35, 1.55
a1 = node(s, "term", A + 0.1, 0.6, W - 0.2, 0.3, "시작 (정류장 조회 요청)")
a2 = node(s, "dec", A, 1.1, W, 0.62, "캐시가 유효한가?\n(30초 이내)")
a3 = node(s, "proc", A, 1.98, W, 0.42, "서울시 도착정보 API 호출\n(노선 전체 · 위치 API)")
a4 = node(s, "dec", A, 2.66, W, 0.62, "응답이 정상인가?")
a5 = node(s, "dec", A, 3.55, W, 0.62, "운행 시간인가?\n(첫차 ~ 막차)")
a6 = node(s, "db", B + 0.15, 1.12, WB - 0.3, 0.52, "조회 캐시\n(30초)", size=7)
x1 = node(s, "err", B, 2.74, WB, 0.46, "API 장애 안내\n마지막 정보 + 갱신 시각")
x2 = node(s, "err", B, 3.63, WB, 0.46, "운행 전 · 종료 안내\n첫차 · 막차 시각 표시")
down(s, a1, a2); down(s, a2, a3, label="N", lpos=(cx(a2) + 0.05, 1.74)); down(s, a3, a4)
down(s, a4, a5, label="Y", lpos=(cx(a4) + 0.05, 3.3))
edge(s, [rgt(a2), lft(a6)], label="Y", lpos=(A + W + 0.02, 1.2))
edge(s, [rgt(a4), lft(x1)], label="N", lpos=(A + W + 0.02, 2.78))
edge(s, [rgt(a5), lft(x2)], label="N", lpos=(A + W + 0.02, 3.67))
b1 = node(s, "dec", Cc, 3.55, WC, 0.62, "첫째 버스 정보가\n있는가?")
b2 = node(s, "proc", Cc, 2.74, WC, 0.46, "버스 정보 정리\n도착 예정 · 남은 정류장 · 재차")
b3 = node(s, "dec", Cc, 1.9, WC, 0.62, "첫째 버스가\n바뀌었는가?")
b4 = node(s, "page", Cc, 1.1, WC, 0.5, "도착정보 화면\n첫째 · 둘째 버스 + 갱신 시각")
b5 = node(s, "term", Cc + 0.3, 0.6, WC - 0.6, 0.3, "끝 (예측 단계로)")
edge(s, [bot(a5), (cx(a5), 4.55), (cx(b1), 4.55), bot(b1)], label="Y", lpos=(cx(a5) + 0.05, 4.2))
edge(s, [top(a6), (cx(a6), 0.98), (cx(b4) - 0.4, 0.98), (cx(b4) - 0.4, 1.1)])
text(s, cx(a6) + 0.05, 0.8, 1.3, 0.2, [("캐시 응답", {"size": 6.5, "color": C["gray"]})])
up(s, b1, b2, label="Y", lpos=(cx(b1) + 0.05, 3.3)); up(s, b2, b3); up(s, b3, b4, label="N", lpos=(cx(b3) + 0.05, 1.66)); up(s, b4, b5)
x3 = node(s, "err", D, 3.63, WD, 0.46, "도착 정보 없음 안내\n(버스 미운행 · 결측)")
x4 = node(s, "proc", D, 1.98, WD, 0.46, "이전 예측 삭제\n새 버스로 재예측")
edge(s, [rgt(b1), lft(x3)], label="N", lpos=(Cc + WC + 0.02, 3.67))
edge(s, [rgt(b3), lft(x4)], label="Y", lpos=(Cc + WC + 0.02, 2.0))
edge(s, [top(x4), (cx(x4), cy(b4)), rgt(b4)])
for x in (x1, x2, x3):
    pass
text(s, B - 0.3, 4.6, 2.3, 0.25, [("장애 · 운행 종료 · 정보 없음은 30초 뒤 자동 재조회", {"size": 6.5, "color": C["gray"]})])
note_card(s, 8.1, 0.6, 1.55, 4.25, "상태별 표시", ic="MdNotificationsActive", bullets=[
    "정상: 도착정보 + 확률",
    "정보 60초 이상 지남: 도착정보만, 확률 보류",
    "API 장애: 마지막 정보 + '갱신 실패' 표시",
    "운행 전 · 종료: 첫차 · 막차 시각",
    "정보 없음: 안내 문구만",
    "API 키는 서버에만 보관 (화면에 노출 안 함)",
], size=7)

# ───────────────────────── 4. 탑승 가능성 예측 · 표시
s = new_slide("기능 흐름도 - 탑승 가능성 예측 · 표시"); banner(s, "대기 인원 입력 · 결과 확인 프로세스", "MdInsights")
points(s, [("MdPeople", "입력 1개", "앞 대기 인원 0~40명"), ("MdInsights", "도착 시 인원 분포", "→ 탈 확률"),
           ("MdVerified", "'여유'는 97%↑", "실제 탑승 검증 구간만"), ("MdSchool", "수업 직전 98%↑", "피크는 더 엄격"),
           ("MdInfoOutline", "% 숫자는 상세에서", "기본은 3단계 표시")])
A, W, B, WB = 0.4, 1.9, 2.55, 1.45
p1 = node(s, "io", A, 0.6, W, 0.38, "앞 대기 인원 N 입력 (F-02)")
p2 = node(s, "dec", A, 1.18, W, 0.6, "0~40명\n정수인가?")
p3 = node(s, "proc", A, 2.0, W, 0.62, "입력값 모으기\n지금 재차 · 남은 정류장 · 정차 중 여부\n시각 · 요일 · 수업 시작 · 휴일", size=7)
p4 = node(s, "dec", A, 2.84, W, 0.6, "버스가 바로\n앞 정류장인가?")
p5 = node(s, "proc", A, 3.7, W, 0.5, "모델 (F-03) → 부록 B\n혼합 규칙 + XGBoost 분위수", size=7)
pe = node(s, "err", B, 1.25, WB, 0.46, "입력 수정 안내")
pk = node(s, "proc", B, 3.7, WB, 0.5, "지금 재차 +\n정류장별 실제 변화 분포", size=7)
pdb = node(s, "db", B + 0.1, 2.03, WB - 0.2, 0.56, "학사일정 · 공휴일\n강의시간표 · 과거 CSV", size=6.5)
down(s, p1, p2); down(s, p2, p3, label="Y", lpos=(cx(p2) + 0.05, 1.8)); down(s, p3, p4)
down(s, p4, p5, label="N", lpos=(cx(p4) + 0.05, 3.46))
edge(s, [rgt(p2), lft(pe)], label="N", lpos=(A + W + 0.02, 1.3))
edge(s, [top(pe), (cx(pe), cy(p1)), rgt(p1)])
edge(s, [rgt(p4), (cx(pk), cy(p4)), top(pk)], label="Y", lpos=(A + W + 0.05, 2.9))
edge(s, [lft(pdb), rgt(p3)], color="blue", dashed=True)
Cc, WC = 4.3, 2.0
q1 = node(s, "proc", Cc - 0.15, 4.35, WC + 0.3, 0.5, "도착 시 재차 분포\nP(탑승) = P(재차 − 하차 ≤ 정원 − N − 1)", size=6.5)
edge(s, [bot(p5), (cx(p5), cy(q1)), lft(q1)])
edge(s, [bot(pk), (cx(pk), cy(q1))], arrow=False)
q2 = node(s, "dec", Cc, 3.55, WC, 0.62, "검증을 통과한\n구간인가? (F-04)")
q3 = node(s, "dec", Cc, 2.72, WC, 0.62, "분포가 너무\n넓은가?")
q4 = node(s, "dec", Cc, 1.89, WC, 0.62, "P ≥ 여유 기준?\n(시간대·거리별)")
q5 = node(s, "dec", Cc, 1.06, WC, 0.62, "P ≤ 0.3 인가?")
up(s, q1, q2); up(s, q2, q3, label="Y", lpos=(cx(q2) + 0.05, 3.35)); up(s, q3, q4, label="N", lpos=(cx(q3) + 0.05, 2.52))
up(s, q4, q5, label="N", lpos=(cx(q4) + 0.05, 1.69))
D, WD = 6.6, 1.35
r0 = node(s, "err", D, 3.63, WD, 0.46, "확률 보류\n도착정보만 표시", size=7)
r1 = node(s, "ok", D, 2.8, WD, 0.46, "판단하기 일러요", size=7.5, bold=True)
r2 = node(s, "ok", D, 1.97, WD, 0.46, "여유 있음", size=8, bold=True)
r3 = node(s, "ok", D, 1.14, WD, 0.46, "어려울 수 있음", size=8, bold=True)
r4 = node(s, "ok", Cc + 0.2, 0.55, WC - 0.4, 0.28, "빠듯함", size=8, bold=True)
for r, col in ((r1, C["gray"]), (r2, C["ok"]), (r3, C["err"]), (r4, C["warn"])):
    r.line.color.rgb = rgb(col)
    for p in r.text_frame.paragraphs:
        for run in p.runs:
            run.font.color.rgb = rgb(col)
edge(s, [rgt(q2), lft(r0)], label="N", lpos=(Cc + WC + 0.02, 3.68))
edge(s, [rgt(q3), lft(r1)], label="Y", lpos=(Cc + WC + 0.02, 2.85))
edge(s, [rgt(q4), lft(r2)], label="Y", lpos=(Cc + WC + 0.02, 2.02))
edge(s, [rgt(q5), lft(r3)], label="Y", lpos=(Cc + WC + 0.02, 1.19))
up(s, q5, r4, label="N", lpos=(cx(q5) + 0.05, 0.86))
res = node(s, "page", 8.1, 1.7, 1.55, 1.0, "결과 화면\n단계 + 근거\n(지금 재차 · 몇 정류장 전 · 갱신 시각)\n둘째 버스 도착", size=6.5)
for r in (r1, r2, r3):
    edge(s, [rgt(r), (8.0, cy(r)), (8.0, cy(res)), lft(res)])
edge(s, [rgt(r4), (8.0, cy(r4)), (8.0, cy(r3))], arrow=False)
edge(s, [rgt(r0), (8.0, cy(r0)), (8.0, cy(r1))], arrow=False)
note_card(s, 8.1, 2.95, 1.55, 1.9, "표시 원칙", ic="MdVerified", bullets=[
    "모델은 정직한 확률, 보수성은 표시 기준에",
    "'여유' 기준은 실제 탑승 97%↑ (수업 직전 98%↑)",
    "% 숫자는 눌러야 보임",
    "정원 = 노선·차종별 max(잠정값, 관측 최대)",
], size=6.5)

# ───────────────────────── 5. 데이터 수집 · 학습 · 평가
s = new_slide("부록 A - 데이터 수집 · 학습 · 평가"); banner(s, "부록 A · 데이터 수집 → 학습 → 평가 (AI · 데이터)", "MdCloudSync")
points(s, [("MdEventAvailable", "버스 시간표대로 수집", "주말 포함 · 새벽 차단"), ("MdTimer", "27 · 40 · 150초", "피크 · 저녁 · 그 외"),
           ("MdStorage", "원본 그대로 보관", "날짜별 JSON"), ("MdVerified", "기준 통과해야 공개", "못 이기면 보류"),
           ("MdSchool", "11주차 독립 평가", "학생 10명 테스트")])
A, W, B, WB = 0.4, 1.85, 2.45, 1.4
c1 = node(s, "term", A + 0.05, 0.6, W - 0.1, 0.3, "시작 (오라클 서버 수집기)")
c2 = node(s, "dec", A, 1.1, W, 0.62, "운행 시간인가?\n(시간표 ±10분)", size=6.8)
c3 = node(s, "proc", A, 1.98, W, 0.5, "버스 API 호출\n2개 노선 · 키 4개 순환")
c4 = node(s, "db", A + 0.15, 2.72, W - 0.3, 0.56, "원본 로그 저장\n(날짜별 JSON)", size=7)
c5 = node(s, "proc", A, 3.52, W, 0.62, "학습 데이터 만들기 (분석 서버)\n도착 판정 → 정답 = 도착 시 재차", size=7)
cw = node(s, "proc", B, 1.18, WB, 0.46, "대기\n(01:30~05:00 호출 금지)", size=7)
cdb = node(s, "db", B + 0.05, 3.55, WB - 0.1, 0.56, "공공 CSV 44개월 · 공휴일\n학사일정 · 강의시간표", size=6.5)
down(s, c1, c2); down(s, c2, c3, label="Y", lpos=(cx(c2) + 0.05, 1.74)); down(s, c3, c4); down(s, c4, c5)
text(s, cx(c4) + 0.05, 3.28, 1.6, 0.22, [("요청 시 분석 서버로 복사", {"size": 6, "color": C["gray"]})])
edge(s, [rgt(c2), lft(cw)], label="N", lpos=(A + W + 0.02, 1.2))
edge(s, [top(cw), (cx(cw), 0.75), rgt(c1)])
text(s, B - 0.1, 1.66, WB + 0.2, 0.6, [("간격 27초 · 40초 · 150초", {"size": 6, "color": C["gray"]}), ("시간표는 요일별 자동 학습", {"size": 6, "color": C["gray"]})], align=PP_ALIGN.CENTER)
edge(s, [lft(cdb), rgt(c5)], color="blue", dashed=True)
Cc, WC = 4.15, 1.95
d1 = node(s, "proc", Cc, 4.2, WC, 0.5, "날짜 단위 분할\n마지막 2주는 봉인 평가용")
d2 = node(s, "proc", Cc, 3.4, WC, 0.52, "모델 학습 · 비교\n규칙 · XGBoost · LightGBM · CatBoost", size=7)
d3 = node(s, "dec", Cc, 2.5, WC, 0.66, "기준 모델 통과?\n('여유' 97%↑)")
d4 = node(s, "proc", Cc, 1.7, WC, 0.5, "모델 고정 · 버전 기록\n(DVC, 10주차)")
d5 = node(s, "proc", Cc, 1.02, WC, 0.44, "서비스 배포 (FastAPI)")
d6 = node(s, "term", Cc + 0.3, 0.55, WC - 0.6, 0.3, "끝")
edge(s, [bot(c5), (cx(c5), 4.85), (cx(d1), 4.85), bot(d1)])
up(s, d1, d2); up(s, d2, d3); up(s, d3, d4, label="Y", lpos=(cx(d3) + 0.05, 2.25)); up(s, d4, d5)
D, WD = 6.35, 1.55
de = node(s, "err", D, 2.6, WD, 0.46, "성능 미달 기록\n확률 공개 보류", size=7)
edge(s, [rgt(d3), lft(de)], label="N", lpos=(Cc + WC + 0.02, 2.62))
edge(s, [bot(de), (cx(de), cy(d2)), rgt(d2)])
text(s, cx(de) - 1.25, 3.17, 1.2, 0.22, [("데이터 더 쌓고 재학습", {"size": 6.5, "color": C["gray"]})])
dv = node(s, "proc", D, 1.0, WD, 0.5, "독립 평가 (F-07)\n11주차 · 학생 10명 테스트", size=7)
dr = node(s, "db", D + 0.1, 1.72, WD - 0.2, 0.5, "탑승 결과 기록\n(F-06 · 익명)", size=6.5)
edge(s, [rgt(d5), lft(dv)])
edge(s, [top(dv), (cx(dv), 0.7), (Cc + WC - 0.3, 0.7)])
edge(s, [top(dr), bot(dv)], color="blue", dashed=True)
note_card(s, 8.1, 0.6, 1.55, 4.25, "검증 규칙 (PRD)", ic="MdVerified", bullets=[
    "조회 시점에 알 수 있는 값만 입력으로",
    "같은 날 · 같은 운행이 학습 / 평가에 섞이지 않게 날짜로 분할",
    "기준(지금 재차 그대로)을 못 이기면 '성능 미달'",
    "재학습하면 새로 평가",
    "누설 점검 · 정답 섞기 대조 통과 (9/24)",
], size=7)


# ───────────────────────── 6. AI 예측 모델 내부 흐름
s = new_slide("부록 B - AI 예측 모델 내부 흐름"); banner(s, "부록 B · 예측 모델 내부 (F-03)", "MdQueryStats")
points(s, [("MdDirectionsBus", "정차 중 보정", "출발 때 갱신되는 재차"), ("MdCloudSync", "오늘 실측 혼합", "휴일 · 방학 대응"),
           ("MdQueryStats", "잔차만 학습", "규칙보다 안 나빠짐"), ("MdTune", "정원으로 제한", "차종별 관측 최대"),
           ("MdInsights", "분포로 출력", "평균이 아니라 범위")])
text(s, 0.35, 0.6, 7.5, 0.25, [("조회 시점의 첫째 버스 1대 → 내 정류장 도착 시 재차 분포 (학습 · 서비스 공통)", {"size": 8, "bold": True, "color": C["navy"]})])
g1 = node(s, "io", 0.3, 0.92, 1.85, 0.56, "조회 시점 스냅샷\n지금 재차 · 남은 정류장\n위치 · 도착 예정", size=6.2)
g2 = node(s, "dec", 2.3, 0.9, 1.7, 0.6, "앞 정류장에\n정차 중인가?", size=7)
g3 = node(s, "proc", 4.25, 0.95, 1.75, 0.5, "미반영 몫 더하기\n(재차는 출발 때 갱신)", size=6.5)
edge(s, [rgt(g1), lft(g2)]); edge(s, [rgt(g2), lft(g3)], label="Y", lpos=(4.02, 1.0))
gdb1 = node(s, "db", 2.35, 1.85, 1.6, 0.52, "과거 CSV 평균\n(학기 · 시간대별)", size=6.5)
gdb2 = node(s, "db", 4.3, 1.85, 1.65, 0.52, "오늘 최근 60분\n실측 (다른 버스)", size=6.5)
g4 = node(s, "proc", 6.25, 1.85, 1.85, 0.52, "남은 구간 순승차 혼합\n(n·오늘 + 3·과거) ÷ (n+3)", size=6.5)
edge(s, [(cx(gdb1), 2.37), (cx(gdb1), 2.52), (6.6, 2.52), (6.6, 2.37)], color="blue", dashed=True)
edge(s, [rgt(gdb2), lft(g4)], color="blue", dashed=True)
g5 = node(s, "proc", 6.25, 2.75, 1.85, 0.5, "규칙 예측\n지금 재차 + 미반영 + 혼합 순승차", size=6.5)
edge(s, [rgt(g3), (7.17, cy(g3)), top(g4)])
edge(s, [(cx(g2), 1.5), (cx(g2), 1.68), (7.17, 1.68)], arrow=False)
text(s, cx(g2) - 0.25, 1.47, 0.3, 0.16, [("N", {"size": 7, "bold": True, "color": C["err"]})])
down(s, g4, g5)
g6 = node(s, "dec", 4.25, 2.7, 1.75, 0.6, "버스가 바로\n앞 정류장인가?", size=7)
edge(s, [lft(g5), rgt(g6)])
g7 = node(s, "proc", 2.3, 2.75, 1.7, 0.5, "지금 재차 그대로 +\n(정류장·정차) 실제 변화 분포", size=6.2)
edge(s, [lft(g6), rgt(g7)], label="Y", lpos=(4.02, 2.8))
g8 = node(s, "proc", 4.25, 3.6, 1.75, 0.55, "XGBoost 분위수 (5~95%)\n규칙이 틀린 만큼(잔차)만 학습", size=6.3)
edge(s, [bot(g6), top(g8)], label="N", lpos=(cx(g6) + 0.05, 3.32))
g9 = node(s, "proc", 6.25, 3.6, 1.85, 0.55, "재차 범위 제한\n0 ≤ 재차 ≤ 정원", size=6.5)
edge(s, [rgt(g8), lft(g9)])
gcap = node(s, "db", 6.35, 4.35, 1.65, 0.52, "정원: 노선·차종별\nmax(잠정 33·55, 관측 최대)", size=6)
edge(s, [top(gcap), bot(g9)], color="blue", dashed=True)
gfeat = node(s, "db", 2.35, 3.6, 1.6, 0.55, "피처: 배차 · 몰림 · 오늘 혼잡\n수업 시작 · 달력 · 위치", size=5.8)
edge(s, [rgt(gfeat), lft(g8)], color="blue", dashed=True)
g10 = node(s, "term", 0.35, 3.6, 1.7, 0.55, "도착 시 재차 분포\n→ 5쪽 P(탑승)", size=6.5)
edge(s, [bot(g9), (cx(g9), 4.28), (1.2, 4.28), bot(g10)])
edge(s, [lft(g7), (1.2, cy(g7)), top(g10)])
note_card(s, 8.3, 0.55, 1.4, 4.35, "왜 이렇게", ic="MdInfoOutline", bullets=[
    "재차는 버스가 정류장을 떠날 때 갱신 → 정차 중이면 그 정류장 승하차가 아직 빠져 있음",
    "과거 평균만 쓰면 휴일에 틀림 → 오늘 실측으로 섞음",
    "모델은 규칙이 틀린 만큼만 배움 → 데이터가 적어도 규칙보다 크게 나빠지지 않음",
    "바로 앞은 변화가 0 아니면 크게 → 평균 대신 분포",
], size=6)

# ───────────────────────── 7. 검증 현황 · 제안
SQ = False; DY = 0.25
s = new_slide("부록 C - 검증 현황 · 제안"); banner(s, "부록 C · 평가 절차 · 중간 결과 · 제안 (F-07)", "MdVerified")
text(s, 0.35, 0.6, 6.0, 0.25, [("평가 절차 (F-07) — 처음 보는 날을 맞히는지로 판정", {"size": 8.5, "bold": True, "color": C["navy"]})])
v1 = node(s, "proc", 0.35, 0.9, 1.45, 0.5, "학습: 앞선 날짜\n평가: 다음 날짜", size=6.8)
v2 = node(s, "proc", 2.0, 0.9, 1.45, 0.5, "기준 모델\n'지금 재차 그대로'", size=6.8)
v3 = node(s, "proc", 3.65, 0.9, 1.45, 0.5, "누설 점검 ·\n정답 섞기 대조", size=6.8)
v4 = node(s, "dec", 5.3, 0.85, 1.5, 0.6, "기준을\n이기는가?", size=7)
v5 = node(s, "ok", 7.0, 0.93, 1.1, 0.44, "구간별 공개", size=7, bold=True)
edge(s, [rgt(v1), lft(v2)]); edge(s, [rgt(v2), lft(v3)]); edge(s, [rgt(v3), lft(v4)])
edge(s, [rgt(v4), lft(v5)], label="Y", lpos=(6.82, 0.93))
v6 = node(s, "err", 5.3, 1.65, 1.5, 0.4, "성능 미달 · 보류", size=7)
edge(s, [bot(v4), top(v6)], label="N", lpos=(cx(v4) + 0.05, 1.45))
text(s, 0.35, 2.2, 6.0, 0.25, [("중간 결과 — 첫 평일 등교일(9/28 월) 예측 · 4일 2,088운행 (9/26~27 미수집)", {"size": 8, "bold": True, "color": C["navy"]})])
stats = [("−25%", "도착 재차 오차", "4.93 → 3.69명"), ("−39%", "확률 오차 (Brier)", "0.079 → 0.048"),
         ("−33%", "3정류장 이상 전 오차", "8.71 → 5.81명"), ("96.7%", "'여유' 표시 → 실제 탑승", "목표 97% 근접 (오후 98.8%)")]
for i, (big, lab, sub) in enumerate(stats):
    x = 0.35 + i * 2.0
    text(s, x, 2.5, 1.85, 0.95, [(big, {"size": 22, "bold": True, "color": C["navy"], "align": PP_ALIGN.CENTER}),
                                  (lab, {"size": 7.5, "bold": True, "align": PP_ALIGN.CENTER}),
                                  (sub, {"size": 7, "color": C["gray"], "align": PP_ALIGN.CENTER})],
         fill=C["note_f"], line="C9D3E6")
text(s, 0.35, 3.55, 7.85, 0.3, [("기준 = 지금 재차 그대로 · 평가 = 9/23~25(평일 오후 + 추석) 학습 → 9/28 하루 예측. 평일 학습 데이터가 쌓이면 재평가. 무작위 도착 때문에 생기는 이론상 최소 오차: 1~2정류장 전 약 2명, 3정류장 이상 전 약 3명", {"size": 6.5, "color": C["gray"]})])
w1 = node(s, "io", 0.35, 4.15, 1.9, 0.55, "(제안) 앞 정류장 학생이\n입력한 대기 인원", size=6.5)
w2 = node(s, "proc", 2.5, 4.15, 2.2, 0.55, "뒤 정류장 예측에 '실제 줄 길이'로\n반영 → 운의 몫을 정보로", size=6.5)
w3 = node(s, "ok", 4.95, 4.2, 1.6, 0.45, "쓸수록 정확해짐", size=7, bold=True)
edge(s, [rgt(w1), lft(w2)]); edge(s, [rgt(w2), lft(w3)])
note_card(s, 8.3, 0.55, 1.4, 4.35, "지금까지 고친 것", ic="MdTune", bullets=[
    "순환노선 집계 오류 (정문 두 줄)",
    "휴일에 과거 평균이 틀림 → 오늘 실측 혼합",
    "정차 중 미반영 승하차 → 위치 API 보정",
    "정원 가정 오류 (종로03 57명) → 관측 최대로",
    "실패가 수업 직전에 몰림 → 시간대별 기준",
    "수집: 주말 포함 · 버스 시간표 기반 · 오라클 이전",
], size=6)

# 원래 5장(기호 설명 · Rento 예시) 제거 → 새 슬라이드만 남김
sld = prs.slides._sldIdLst
for sid in list(sld)[:5]:
    prs.part.drop_rel(sid.get(qn("r:id")))
    sld.remove(sid)

NOTES = [
"""[읽는 법] 1쪽 기호 → 2쪽 전체 단계 요약 → 3쪽 전체 상세 → 4~5쪽 기능별 흐름. 6~8쪽은 AI·데이터 쪽 부록(수집·모델·검증)이에요.
[범위] 대상은 등교 방향 6개 정류장(성북02 3곳 + 종로03 3곳)이에요. PRD·WBS 원문은 '등·하교, 양쪽 정류장'이라 범위 변경은 팀 합의가 필요해요.
[기능 번호] PRD의 F-01~F-07을 쪽마다 표시했어요. 쪽 번호 표로 어느 기능이 어디 있는지 찾으면 돼요.""",
"""[전체 요약] 템플릿 예시(전체) 형식으로 7단계를 한 줄에 놓았어요. 남색 = 정보 확인, 초록 = 예측·판단, 회색 = 갱신·기록, 주황 점선 = 예외 처리.
[대기 여부 판단] PRD 대로 확률을 보고 기다릴지 다른 방법을 쓸지는 사용자가 정해요. 서비스는 판단 근거만 줘요.
[공통 지원 기능] 모든 단계에 걸쳐 동작하는 기능이에요 — 자동 갱신, 상태 안내, 검증된 확률만 표시, 개인정보 미저장, (선택) 평가 기록.""",
"""[전체 흐름] 접속 → 정류장 선택 → 도착정보 조회 → 앞 대기 인원 입력 → 예측 → 표시 조건 판단 → 결과 → 재조회 → (선택) 탑승 결과 기록.
[FE] 필요한 화면: 정류장 선택, 결과(여유/빠듯/어려움 + 근거 + 둘째 버스 + 갱신 시각), 상태 안내(운행 종료·장애·정보 없음), 입력 오류 안내.
[BE] 장애·운행 종료일 때 30초 뒤 자동 재조회, 대상 버스가 바뀌면 이전 예측 삭제.
[합의 필요] 30초 재조회·60초 유효시간은 제안값이에요. F-06 탑승 결과 기록은 '선택'으로 그렸어요.""",
"""[BE 담당] 버스 API 는 서버에서만 호출하고 키는 화면에 노출하지 않아요(PRD 원칙 5).
[캐시] 같은 정류장 조회는 30초 안이면 캐시로 응답해 API 한도를 아껴요.
[상태 5종] 정상 / 정보 60초 이상 지남 / API 장애 / 운행 전·종료 / 정보 없음 → FE 상태 화면(WBS 2.11)과 1:1로 맞추면 돼요.
[운행 시간] 버스 도착정보 API 응답의 첫차·막차 시각(firstTm, lastTm)을 그대로 쓰면 돼요.""",
"""[입력] 앞 대기 인원은 0~40명 정수만 받아요.
[결과 5가지] 여유 있음 / 빠듯함 / 어려울 수 있음 / 판단하기 일러요(불확실할 때) / 확률 보류(검증 안 된 구간).
[표시 원칙] 모델은 정직한 확률을 내고, 보수성은 표시 기준에만 넣어요. '여유'는 실제로 97% 이상 탔던 구간에서만 띄우고, 수업 직전 피크는 98%로 더 엄격해요.
[FE] % 숫자는 기본 화면에 내세우지 말고 눌러야 보이게 해주세요. 하단 안내 문구 예시: '실시간 데이터 기반 추정이며 현장 상황에 따라 다를 수 있습니다.'""",
"""[수집] 오라클 서버에서 수집기가 계속 돌아요(9/28부터, 재부팅돼도 자동 재시작). 버스 시간표대로 첫차 10분 전~막차 10분 후에만 호출하고 새벽 01:30~05:00 은 절대 호출하지 않아요.
[학습 데이터] 버스가 정류장에 도착한 순간의 재차인원을 정답으로 자동으로 만들어서, 사람이 직접 관측할 필요가 없어요(WBS 1.5·2.2 '정규 관측' 대체 제안).
[평가] 날짜 단위로 나누고, 기준(지금 재차 그대로)을 못 이기면 '성능 미달'로 기록하고 확률을 공개하지 않아요(PRD 원칙 3).
[일정] 10주차 모델 고정, 11주차 독립 평가 + 학생 10명 테스트(마일스톤과 동일).""",
"""[한 줄 요약] 지금 버스에 타 있는 인원에서 출발해서, 도착 전까지 더 타고 내릴 인원을 더해 '도착 시 인원'의 분포를 구해요.
[포인트 1] 재차 인원은 버스가 정류장을 떠날 때 갱신돼요. 그래서 버스가 앞 정류장에 서 있으면 그 정류장 승하차를 따로 더해요.
[포인트 2] 과거 평균만 쓰면 휴일에 크게 틀려서(추석에 확인), 오늘 실측이 쌓일수록 오늘 흐름을 따라가게 섞어요.
[포인트 3] XGBoost 는 규칙이 틀린 만큼만 배우게 해서, 데이터가 적어도 규칙보다 크게 나빠지지 않아요.
[정원] 노선·차종별로 잠정값과 실제 관측 최대 중 큰 값을 써요(종로03 은 57명까지 관측).""",
"""[숫자 조건] 9/23~25(평일 오후 + 추석)로 학습하고 9/28(월) 첫 평일 등교일 하루를 예측한 결과예요. 9/26~27은 수집 설정 문제로 빠졌어요.
[의미] 기준(지금 재차 그대로)보다 오차 25%, 확률 오차 39% 줄었어요. 멀리 있는 버스일수록 차이가 커요.
['여유' 96.7%] 목표 97%에 조금 못 미쳐요. 평일 학습 데이터가 거의 없어서예요. 평일 데이터가 쌓이면 다시 평가해요(오늘 오후만 보면 98.8%).
[제안] 앞 정류장 학생이 입력한 대기 인원을 뒤 정류장 예측에 쓰면, 운에 맡기던 부분을 실제 정보로 바꿀 수 있어요. 서비스 설계에 넣을지 팀 논의가 필요해요.""",
]
for sl, n in zip(prs.slides, NOTES):
    sl.notes_slide.notes_text_frame.text = n
prs.save(OUT)
print("saved", OUT, "slides", len(prs.slides))
